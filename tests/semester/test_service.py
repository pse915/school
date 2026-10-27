from copy import deepcopy
from datetime import date
from threading import Event
import json,time
import pytest
from tests.operations.fixtures import integrated_app
from tests.fixtures import command
from backend.semester.install import get_service,TAB
from backend.semester.model import compile_project,digest,validate,SemesterError,RevisionConflict
from backend.semester.solver import solve

@pytest.fixture
def env(tmp_path):
    app,ops,legacy,source,sheets=integrated_app(tmp_path/'semester.sqlite3')
    svc=get_service(app);tid=svc.create(2026,1,'학기 테스트','2026-10-12','2026-12-31')
    raw=svc.get(tid)['project'];raw.update(periods_by_day={'월':4,'화':4,'수':4,'목':4,'금':4},
        teachers=[{'id':'T1','name':'테스트가','subject':'과학','max_weekly':20,'blocked':[]},{'id':'T2','name':'테스트나','subject':'수학','max_weekly':20,'blocked':[]}],
        classes=[{'id':'C1','name':'1-1','grade':1,'size':2,'blocked':[]},{'id':'C2','name':'2-1','grade':2,'size':2,'blocked':[]}],
        rooms=[{'id':'R1','name':'과학실','capacity':10,'blocked':[]},{'id':'R2','name':'수학실','capacity':10,'blocked':[]}],
        subjects=[{'id':'S1','name':'과학','credits':3},{'id':'S2','name':'수학','credits':3}],
        students=[{'id':'A','name':'학생A','class_id':'C1'},{'id':'B','name':'학생B','class_id':'C2'}],
        curriculum=[{'id':'L1','subject_id':'S1','teacher_ids':['T1'],'class_ids':['C1'],'student_ids':[],'room_ids':['R1'],'weekly_periods':3,'split':[2,1],'fixed_slots':[],'forbidden_slots':[]},
                    {'id':'L2','subject_id':'S2','teacher_ids':['T2'],'class_ids':['C1'],'student_ids':[],'room_ids':['R2'],'weekly_periods':2,'split':[1,1],'fixed_slots':[],'forbidden_slots':[]}])
    raw['weights']={k:0 for k in raw['weights']};svc.save(tid,raw,1)
    app.port.session_state['semester_term_id']=tid
    return app,svc,tid,raw,ops

def candidate(env):
    app,svc,tid,raw,ops=env;p=compile_project(raw);r=solve(p,seconds=2)
    assert r['status'] in ('FEASIBLE','OPTIMAL'),r
    ident=svc.add_candidate(tid,2,digest(raw),p,r);return ident,p,r

def test_split_credits_and_shared_teachers(env):
    _,svc,tid,raw,_=env;p=compile_project(raw)
    assert [x['duration'] for x in p['lessons'] if x['source_request_id']=='L1']==[2,1]
    assert raw['subjects'][0]['credits']==3
    raw['curriculum'][0]['parts']=[{'duration':2,'teacher_ids':['T1']},{'duration':1,'teacher_ids':['T2']}]
    assert compile_project(raw)['lessons'][1]['teacher_ids']==['T2']
    raw['curriculum'][0]['parts'][0]['teacher_ids']=['T1','T2']
    assert compile_project(raw)['lessons'][0]['teacher_ids']==['T1','T2']
    raw['curriculum'][0]['weekly_periods']=4
    with pytest.raises(SemesterError,match='분할'):compile_project(raw)

def test_scope_copy_and_concurrent_input_changes(env):
    app,svc,tid,raw,ops=env
    copied=svc.create(2027,2,'복사학기','2027-08-01','2027-12-31',tid)
    p=svc.get(copied)['project'];assert p['teachers']==raw['teachers'] and not p['students'] and not p['credit_links'] and not p['curriculum']
    svc.save(tid,raw,2)
    with pytest.raises(RevisionConflict):svc.save(tid,raw,2)
    svc.school='other-school'
    with pytest.raises(SemesterError):svc.get(tid)

def test_solver_and_independent_revalidation(env):
    cid,p,r=candidate(env);assert not validate(p,r['placements'])
    broken=deepcopy(r['placements']);broken[0]['start']='금99'
    assert validate(p,broken)
    with pytest.raises(SemesterError):env[1].add_candidate(env[2],2,digest(env[3]),p,{**r,'placements':broken})

def test_actual_clock_detects_different_numbered_period_overlap(env):
    raw=deepcopy(env[3]);raw['curriculum'][1]['class_ids']=['C2'];raw['curriculum'][1]['teacher_ids']=['T1']
    raw['curriculum']=[{**r,'weekly_periods':1,'split':[1],'room_ids':[]} for r in raw['curriculum']]
    raw['bell_schedules']={'1':{'1':['09:00','09:45']},'2':{'2':['09:30','10:15']}}
    p=compile_project(raw);errors=validate(p,[{'session_id':'L1__1#1','start':'월1','room_id':''},{'session_id':'L2__1#1','start':'월2','room_id':''}])
    assert any('실제 시각 중복' in e for e in errors)

def test_impossible_unknown_cancelled_are_distinct(env):
    p=compile_project(env[3]);p['teachers'][0]['blocked']=[d+str(n) for d in '월화수목금' for n in range(1,5)]
    assert solve(p,1)['status']=='INVALID'
    p=compile_project(env[3]);event=Event();event.set();assert solve(p,1,cancel=event)['status']=='CANCELLED'

def test_finalize_stale_and_idempotent_receipt(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env
    ident=svc.finalize(tid,cid,'검토 완료','stable-key');assert svc.finalize(tid,cid,'검토 완료','stable-key')==ident
    svc.save(tid,raw,2)
    with pytest.raises(RevisionConflict):svc.finalize(tid,cid,'새로운 확정','different-key')

def test_atomic_apply_pointer_phase6_and_legacy_read(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env;before=app.port.session_state['timetable'].copy(deep=True)
    final=svc.finalize(tid,cid,'검증 완료','final-key');svc.apply(tid,final,'2026-10-12')
    assert svc.apply(tid,final,'2026-10-12',0)==1
    from backend.semester.integration import active
    project=active(app,'2026-10-12');assert project['semester_version_id']==final
    from backend.operations.source import build_project
    assert build_project(app,2026,1)['semester_version_id']==final
    daily=ops.day(project,'2026-10-12',ops.actor());assert daily['base_id'].startswith('semester-')
    frame=app.program.call('get_effective_timetable_for_date','2026-10-12',0)
    assert set(frame['교사명'])<= {'테스트가','테스트나'}
    assert app.port.session_state['timetable'].equals(before)

def test_future_changes_block_apply_without_partial_pointer(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env
    legacy=__import__('backend.operations.source',fromlist=['build_project']).build_project(app,2026,1)
    lesson=ops.day(legacy,'2026-10-12',ops.actor())['lessons'][0]
    req=ops.submit(legacy,ops.actor(),'move',[{'action':'move','date':'2026-10-12','occurrence_id':lesson['id'],'to_period':4}],'미래 운영 변경')
    final=svc.finalize(tid,cid,'검증 완료','final-key')
    with pytest.raises(SemesterError,match='검토 필요'):svc.apply(tid,final,'2026-10-12')
    assert not svc.pointer(2026,1,'2026-10-12')
    assert ops.requests(legacy,ops.actor())[0]['request_id']==req

def test_new_menu_all_ten_pages_render(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env
    command(app,'navigate',{'tab':TAB})
    for index in range(10):
        result=command(app,'control',{'control':'semester_page','values':{'semester_page':index}})
        assert result['response']['success'],(index,result['response'])
        assert not any(n['kind']=='error' for n in app._walk(result['nodes'])),(index,result['nodes'])

def test_excel_import_and_exports_and_no_raw_student_ids_in_editor(env):
    from backend.semester.io import template,parse_rows,export_excel,export_html
    from openpyxl import load_workbook
    from io import BytesIO
    cid,p,r=candidate(env)
    assert set(load_workbook(BytesIO(template())).sheetnames)>={'teachers','curriculum'}
    with pytest.raises(SemesterError):parse_rows([{'id':'x'},{'id':'x'}],'teachers')
    assert '학생별' in load_workbook(BytesIO(export_excel(p,r))).sheetnames
    assert b'<!doctype html>' in export_html(p,r)
    app,svc,tid,raw,ops=env;command(app,'navigate',{'tab':TAB});result=command(app,'control',{'control':'semester_page','values':{'semester_page':6}})
    node=next(n for n in app._walk(result['nodes']) if n['kind']=='semester_editor')
    assert 'student_ids' not in json.dumps(node)

def test_background_jobs_discard_stale_results_and_cancel(env):
    app,svc,tid,raw,ops=env
    from backend.semester.jobs import start,poll,stop
    start(app,svc,tid,1,1,1);svc.save(tid,raw,2)
    app._semester_job['future'].result(timeout=10)
    assert poll(app,svc)['status']=='ERROR' and not svc.candidates(tid)
    start(app,svc,tid,1,1,1);stop(app);assert poll(app,svc)['status']=='CANCELLED'

def test_true_infeasible_and_timeout_statuses(env):
    raw=deepcopy(env[3]);raw['curriculum']=[{**r,'weekly_periods':1,'split':[1],'teacher_ids':['T1'],'fixed_slots':['월1'],'room_ids':[]} for r in raw['curriculum']]
    assert solve(compile_project(raw),1)['status']=='INFEASIBLE'
    assert solve(compile_project(env[3]),0)['status']=='UNKNOWN'

def test_manual_move_lock_undo_parent_and_rollback(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env;source=r['placements'][0]
    changed=None
    for day in '월화수목금':
        for n in range(1,4):
            if day+str(n)==source['start']:continue
            try:changed=svc.edit(tid,cid,[{'session_id':source['session_id'],'start':day+str(n)}]);break
            except SemesterError:pass
        if changed:break
    assert changed and svc.candidate(tid,changed)['parent_id']==cid
    locked=svc.edit(tid,changed,[],(source['session_id'],True));item=next(x for x in svc.candidate(tid,locked)['result']['placements'] if x['session_id']==source['session_id'])
    with pytest.raises(SemesterError):svc.edit(tid,locked,[{'session_id':source['session_id'],'start':('화' if item['start'][0]!='화' else '수')+'1'}])
    first=svc.finalize(tid,cid,'초기 확정','f1');second=svc.finalize(tid,changed,'조정 확정','f2')
    assert svc.apply(tid,first,'2026-10-12')==1
    assert svc.apply(tid,second,'2026-10-12',1)==2
    assert svc.apply(tid,first,'2026-10-12',2)==3
    from backend.semester.integration import active
    assert active(app,'2026-10-12')['published']['placements']==r['placements']

def test_apply_failure_rolls_back_base_and_pointer(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env;final=svc.finalize(tid,cid,'초기 확정','receipt')
    with svc.repo.tx(True) as db:
        before=db.execute('SELECT count(*) FROM base_versions').fetchone()[0]
        db.execute("CREATE TRIGGER fail_pointer BEFORE INSERT ON semester_pointers BEGIN SELECT RAISE(ABORT,'simulated pointer failure');END")
    import sqlite3
    with pytest.raises(sqlite3.IntegrityError):svc.apply(tid,final,'2026-10-12')
    with svc.repo.tx() as db:assert db.execute('SELECT count(*) FROM base_versions').fetchone()[0]==before
    assert svc.pointer(2026,1,'2026-10-12') is None

def test_current_semester_permissions_are_not_forged_in_payload(env):
    app,svc,tid,raw,ops=env;app.port.session_state.user_role='일반교사'
    with pytest.raises(PermissionError):svc.save(tid,raw,2)
    with pytest.raises(PermissionError):svc.finalize(tid,'fake','권한 확인','key')

def test_credit_snapshot_mapping_no_omission_and_scope_rejection(env):
    from datetime import datetime,timedelta,timezone
    from diple.sectioning import generate_plan,SectioningPlanService
    app,svc,tid,raw,ops=env;credit=app._credit_service;n=datetime.now(timezone.utc);school='qa-school';token=credit.admin_token
    cid=credit.create_campaign(token,school,{'academic_year':2026,'semester':1,'grade':1,'manager':'tester','title':'과목선택','starts_at':(n-timedelta(days=1)).isoformat(),'ends_at':(n+timedelta(days=1)).isoformat(),'edits_until':(n+timedelta(days=2)).isoformat()})
    group=credit.add_group(token,school,cid,'선택',1,1,target_units=1);off=credit.add_offering(token,school,cid,group,course_id='choice',course_name='선택과학',credits=3,units=1)
    codes=credit.import_roster(token,school,2026,1,[{'student_id':'Z','grade':1,'class_id':'1-1','number':1,'name':'합성학생'}],{'1-1':1})
    credit.activate(school,'Z',codes[0]['activation_code'],'synthetic_password_123');student=credit.login(school,'Z','synthetic_password_123');credit.publish(token,school,cid);credit.submit(student,school,cid,[off],'stable-choice-key',0)
    credit.advance(token,school,cid,'신청 마감');credit.advance(token,school,cid,'검토 중');snap=credit.finalize(token,school,cid)
    plan=SectioningPlanService(credit).save(token,school,snap,generate_plan(snap),ack_warnings=True)
    sid=plan['sections'][0]['section_id']
    revision=svc.link_credit(tid,2,plan['plan_id'],{sid:'T1'},{sid:'R1'},keep_slots=False,class_map={'1-1':'C1'})
    updated=svc.get(tid)['project'];assert revision==3 and any(r.get('student_ids')==['Z'] for r in updated['curriculum'])
    with credit.repo.tx() as db:stored=json.loads(db.execute('SELECT snapshot_json FROM snapshots WHERE snapshot_id=?',(snap['snapshot_id'],)).fetchone()[0])
    assert stored==snap

def test_expired_semester_is_not_used_for_new_daily_states(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env;final=svc.finalize(tid,cid,'확정 검증','expiry')
    svc.apply(tid,final,'2026-10-12')
    from backend.semester.integration import active
    assert active(app,'2027-01-01') is None
    with ops._db() as db:
        base=ops._base(db,('qa-school',2026,1),'2027-01-01')
        assert not json.loads(base['project_json']).get('semester_version_id')

def test_duplicate_jobs_across_sessions_and_abandoned_lease_are_rejected(env):
    from tests.operations.fixtures import integrated_app
    from backend.semester.jobs import start,poll,stop
    from backend.semester.install import get_service
    app,svc,tid,raw,ops=env
    path=svc.repo.path
    other,*_=integrated_app(path);other_svc=get_service(other)
    start(app,svc,tid,1,1,1)
    with pytest.raises(SemesterError,match='다른 세션'):start(other,other_svc,tid,1,2,1)
    stop(app);assert poll(app,svc)['status']=='CANCELLED'
    start(other,other_svc,tid,1,2,1);other._semester_job['future'].result(timeout=10)
    with svc.repo.tx(True) as db:db.execute('DELETE FROM semester_job_locks WHERE term_id=?',(tid,))
    assert poll(other,other_svc)['status']=='ERROR' and not svc.candidates(tid)

def test_old_dated_editor_cannot_edit_a_new_semester_base(env):
    cid,p,r=candidate(env);app,svc,tid,raw,ops=env;final=svc.finalize(tid,cid,'확정 승인','base-stale')
    svc.apply(tid,final,'2026-10-12')
    from backend.semester.integration import active
    project=active(app,'2026-10-12');state=ops.day(project,'2026-10-12',ops.actor());lesson=state['lessons'][0]
    result=command(app,'ops.preview',{'day':'2026-10-12','revision':state['rev'],'base_id':'old-version','occurrence_id':lesson['id'],'to_period':4,'reason':'이전 화면 요청'})
    assert not result['response']['success'] and '버전' in result['response']['error']['message']

def test_reference_candidate_minimal_change_and_fixed_subset(env):
    app,svc,tid,raw,ops=env;cid,p,r=candidate(env)
    ref={x['session_id']:{k:x.get(k,'') for k in ('start','room_id')} for x in r['placements']}
    p['weights']={k:0 for k in p['weights']};p['weights']['changes']=100
    result=solve(p,1,reference=ref)
    assert result['status'] in ('FEASIBLE','OPTIMAL') and result['metrics']['changes']==0
    from backend.semester.jobs import start,poll
    first=r['placements'][0];start(app,svc,tid,1,1,1,reference=r,fixed_ids=[first['session_id']])
    app._semester_job['future'].result(timeout=10);saved=poll(app,svc)
    assert saved['candidate_ids']
    new=svc.candidate(tid,saved['candidate_ids'][0]);assert next(x['start'] for x in new['result']['placements'] if x['session_id']==first['session_id'])==first['start']

def test_input_revision_history_preserves_before_and_after(env):
    app,svc,tid,raw,ops=env;updated=deepcopy(raw);updated['term']['name']='수정 학기'
    svc.save(tid,updated,2)
    with svc.repo.tx() as db:
        history=[dict(r) for r in db.execute('SELECT * FROM semester_term_history WHERE term_id=? ORDER BY revision',(tid,))]
    assert [r['revision'] for r in history]==[1,2,3]
    assert json.loads(history[1]['project_json'])['term']['name']=='학기 테스트'
    assert json.loads(history[2]['project_json'])['term']['name']=='수정 학기'

def test_unknown_class_or_room_capacity_cannot_be_finalized_as_safe(env):
    raw=deepcopy(env[3]);raw['classes'][0]['size']=0;raw['students']=[]
    with pytest.raises(SemesterError,match='학급 인원'):compile_project(raw)
    raw=deepcopy(env[3]);raw['rooms'][0]['capacity']=0
    with pytest.raises(SemesterError,match='수용인원'):compile_project(raw)
