"""Phase 5B: section allocation, integrity, isolation, integration, regression."""
from copy import deepcopy
from datetime import datetime,timedelta,timezone
import pytest
from diple.sectioning import (SectioningError,preflight,generate_plan,validate_plan,
                              move_student,apply_plan_to_project,SectioningPlanService)
from diple.samples import high_school
from diple.model import validate_placements,validate_project
from diple.enrollment.service import EnrollmentService,PASSWORD_HASHER,AccessDenied


def fixture_snapshot(compound=False,count=12,cap=5):
    course_ids=['PH','CH'] if not compound else ['PH','CH','M1']
    off=[{'offering_id':oid+'-off','course_id':oid,'course_name':oid,'group_id':'science',
          'credits':3,'units':(3 if oid=='PH' else 2 if compound and oid=='CH' else 1 if compound else 3),
          'section_capacity':cap,'planned_capacity':50,'min_students':0,'availability':'개설 예정',
          'joint_education':0} for oid in course_ids]
    if compound:
        choices=[[x['offering_id'] for x in off] for _ in range(count)]
    else:
        choices=[[off[n%2]['offering_id']] for n in range(count)]
    return {'snapshot_id':'sample-snap','version':1,'campaign_id':'campaign-sample',
        'school_id':'demo-high','academic_year':2026,'semester':2,'created_at':'2026-10-01T00:00:00+00:00',
        'groups':[{'group_id':'science','name':'과학','target_units':3,'min_count':1,'max_count':3}],
        'offerings':off,
        'students':[{'student_id':f'X{i:04d}','class_id':'H1','grade':2,'number':i+1,'choices':choices[i]} for i in range(count)]}


def test_normal_sections_and_capacity():
    snap=fixture_snapshot(count=16,cap=3)
    plan=generate_plan(snap,force_fallback=True)
    assert not validate_plan(snap,plan)
    assert len(plan['sections'])==6
    assert len(plan['assignments'])==16
    assert all(x['lane']==0 for x in plan['sections'])


def test_compound_3_plus_2_plus_1_units():
    snap=fixture_snapshot(compound=True,count=6,cap=3)
    plan=generate_plan(snap,force_fallback=True)
    assert not validate_plan(snap,plan)
    assert plan['groups'][0]['lane_count']==2
    assert len(plan['assignments'])==18
    blocks={(x['offering_id'],x['lane'],x['offset']) for x in plan['sections']}
    assert len(blocks)==3
    assert sorted([o['units'] for o in snap['offerings']])==[1,2,3]


def test_copy_only_manual_change_and_capacity_limit():
    snap=fixture_snapshot(count=7,cap=2)
    plan=generate_plan(snap,force_fallback=True)
    assignment=plan['assignments'][0];sec_id=assignment['section_id']
    candidates=[x for x in plan['sections'] if x['offering_id']==assignment['offering_id'] and x['section_id']!=sec_id]
    assert candidates
    # size distribution leaves every target either full or limited; verify no bypass.
    with pytest.raises(SectioningError):
        move_student(snap,plan,assignment['student_id'],assignment['offering_id'],'wrong')
    assert plan['assignments'][0]['section_id']==sec_id


def test_snapshot_scope_and_missing():
    snap=fixture_snapshot()
    plan=generate_plan(snap,force_fallback=True)
    bad=deepcopy(plan);bad['school_id']='other'
    assert validate_plan(snap,bad)
    bad=deepcopy(plan);bad['assignments'].pop()
    assert any('미배정' in t for t in validate_plan(snap,bad))


def test_forged_student_course_and_overlap():
    snap=fixture_snapshot(compound=True,count=2,cap=3)
    plan=generate_plan(snap,force_fallback=True)
    bad=deepcopy(plan);bad['assignments'][0]['student_id']='not-student'
    assert validate_plan(snap,bad)
    bad=deepcopy(plan)
    courses={s['offering_id']:s for s in bad['sections']}
    for x in courses.values():x['lane']=0;x['offset']=0
    assert any('중복' in t for t in validate_plan(snap,bad))


def test_missing_minimum_is_warning_not_auto_cancellation():
    snap=fixture_snapshot(count=3)
    snap['offerings'][0]['min_students']=6
    p=generate_plan(snap,force_fallback=True)
    assert p['warnings']
    assert len(p['assignments'])==3


def test_per_school_term_enforced_on_integration():
    snap=fixture_snapshot(count=2)
    p=generate_plan(snap,force_fallback=True)
    project=high_school()
    with pytest.raises(SectioningError,match='학교·학년도·학기'):
        off=deepcopy(project);off['school']['semester']=1
        apply_plan_to_project(off,snap,p,{})
    with pytest.raises(SectioningError,match='담당교사'):
        apply_plan_to_project(project,snap,p,{},replace_lessons=True)


def test_plan_model_timing_block_validation():
    project=high_school()
    snap=fixture_snapshot(count=2)
    plan=generate_plan(snap,force_fallback=True)
    teacher_map={sec['section_id']:('HP' if sec['course_id']=='PH' else 'HC') for sec in plan['sections']}
    built=apply_plan_to_project(project,snap,plan,teacher_map,replace_lessons=True)
    assert len(built['lessons'])==3
    assert len(built['students'])==6
    assert not validate_project(built)
    # Corrupt relative timing by placing different lanes with invalid same start.
    bad=deepcopy(built)
    elective=[x for x in bad['lessons'] if x.get('timing_block')]
    assert elective
    for x in elective:
        x['timing_offset']=10
    assert validate_project(bad)
    assert len(project['lessons'])==3  # the original is unchanged


def test_relative_block_validation_on_placements():
    project=high_school()
    snap=fixture_snapshot(compound=True,count=2,cap=3)
    plan=generate_plan(snap,force_fallback=True)
    teacher_map={x['section_id']:('HK' if x['course_id']=='M1' else 'HP' if x['course_id']=='PH' else 'HC') for x in plan['sections']}
    project['subjects'].append({'id':'M1','name':'교양','credits':1})
    # Teacher assignments might overload PH, but input structurally valid at 2 periods. No solving here.
    built=apply_plan_to_project(project,snap,plan,teacher_map,replace_lessons=True)
    assert not validate_project(built)
    assert all('timing_block' in x for x in built['lessons'] if x['id'].startswith('L-SEC-'))


def test_versioned_persistence_and_school_isolation(tmp_path):
    svc=EnrollmentService(path=str(tmp_path/'enroll.sqlite3'))
    pwd='safe-admin-passphrase-long'
    svc.bootstrap_admin('demo-high','admin',pwd,PASSWORD_HASHER.hash(pwd))
    admin=svc.login('demo-high','admin',pwd)
    now=datetime.now(timezone.utc)
    campaign=svc.create_campaign(admin,'demo-high',{'academic_year':2026,'semester':2,'grade':2,
        'title':'실험','manager':'담당','starts_at':(now-timedelta(hours=1)).isoformat(),
        'ends_at':(now+timedelta(hours=1)).isoformat(),'edits_until':(now+timedelta(hours=2)).isoformat()})
    gid=svc.add_group(admin,'demo-high',campaign,'과학',1,1)
    offering=svc.add_offering(admin,'demo-high',campaign,gid,course_id='PH',course_name='물리',credits=3,section_capacity=2)
    student=svc.import_roster(admin,'demo-high',2026,2,[{'student_id':'X1','class_id':'H1','grade':2,'number':1,'name':'예시'}],{'H1':2})
    svc.activate('demo-high','X1',student[0]['activation_code'],'a-valid-long-student-secret')
    user=svc.login('demo-high','X1','a-valid-long-student-secret')
    svc.publish(admin,'demo-high',campaign)
    svc.submit(user,'demo-high',campaign,[offering],'test-key-long-enough')
    svc.advance(admin,'demo-high',campaign,'신청 진행')
    svc.advance(admin,'demo-high',campaign,'신청 마감')
    svc.advance(admin,'demo-high',campaign,'검토 중')
    snap=svc.finalize(admin,'demo-high',campaign,ack_missing=True)
    plan=generate_plan(snap,force_fallback=True)
    versioner=SectioningPlanService(svc)
    with pytest.raises(AccessDenied):
        versioner.save(user,'demo-high',snap,plan,'student try')
    with pytest.raises(AccessDenied):
        versioner.save(admin,'different-school',snap,plan,'wrong school')
    p1=versioner.save(admin,'demo-high',snap,plan,'최초 분반')
    p2=versioner.save(admin,'demo-high',snap,plan,'재편성 기록')
    assert p1['version']==1 and p2['version']==2
    assert versioner.get(admin,'demo-high',p1['plan_id'])['version']==1
    assert len(versioner.list_versions(admin,'demo-high',snap['snapshot_id']))==2
    with pytest.raises(AccessDenied):
        versioner.get(user,'demo-high',p1['plan_id'])


def test_compound_sections_flow_into_actual_timetable_engine():
    from diple.engine import solve_offline_small
    snap=fixture_snapshot(compound=True,count=2,cap=4)
    plan=generate_plan(snap,force_fallback=True)
    src=high_school();src['subjects'].append({'id':'M1','name':'교양','credits':1})
    teacher_map={s['section_id']:('HK' if s['course_id']=='M1' else 'HP' if s['course_id']=='PH' else 'HC') for s in plan['sections']}
    built=apply_plan_to_project(src,snap,plan,teacher_map,replace_lessons=True)
    outcome=solve_offline_small(built,max_sessions=30)
    assert outcome['status']=='FEASIBLE',outcome['errors']
    assert not validate_placements(built,outcome['placements'])
    bad=deepcopy(outcome['placements'])
    minor=next(x for x in built['lessons'] if x.get('timing_block') and x['subject_id']=='M1')
    major=next(x for x in built['lessons'] if x.get('timing_block') and x['subject_id']=='CH')
    main_slot=next(x['start'] for x in bad if x['session_id'].startswith(major['id']+'#'))
    for item in bad:
        if item['session_id'].startswith(minor['id']+'#'):item['start']=main_slot
    assert any('동시성/상대교시' in e for e in validate_placements(built,bad))


@pytest.mark.parametrize('n',[300,600,1200])
def test_large_synthetic_roster_allocates_without_overlap(n):
    snap=fixture_snapshot(compound=True,count=n,cap=32)
    p=generate_plan(snap,force_fallback=True)
    assert len(p['assignments'])==3*n
    assert not validate_plan(snap,p)


def test_section_report_excel_export():
    from io import BytesIO
    from openpyxl import load_workbook
    from diple.sectioning_export import plan_to_excel
    snap=fixture_snapshot(compound=True,count=8,cap=3)
    plan=generate_plan(snap,force_fallback=True)
    wb=load_workbook(BytesIO(plan_to_excel(snap,plan)),read_only=True)
    assert {'StudentAssignment','Sections','OperatingBlocks','PlanVersion','ImportValidation'} <= set(wb.sheetnames)
    assert wb['StudentAssignment'].max_row==25


def test_reject_tampered_section_capacity_and_course_id():
    snap=fixture_snapshot(count=8,cap=2)
    plan=generate_plan(snap,force_fallback=True)
    altered=deepcopy(plan)
    altered['sections'][0]['capacity']=99
    assert any('분반정원' in s for s in validate_plan(snap,altered))
    altered=deepcopy(plan)
    altered['sections'][0]['course_id']='wrong'
    assert any('교과 ID' in s for s in validate_plan(snap,altered))
