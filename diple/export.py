"""Excel export, not a certified NEIS import file."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .model import index, expand_sessions, DAYS


def excel_safe(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def timetable_rows(project, placements, perspective, ident):
    lesson_idx = {x['session_id']: x for x in expand_sessions(project)}
    names = {k:index(project,k) for k in ('teachers','classes','rooms','subjects','students')}
    cells = {}
    for p in placements:
        s = lesson_idx[p['session_id']]
        if perspective == 'rooms':
            included = (p.get('room_id','') == ident)
        elif perspective == 'students':
            student = names['students'].get(ident,{})
            included = ident in s.get('student_ids',[]) or student.get('class_id') in s.get('class_ids',[])
        else:
            included = ident in s.get('teacher_ids' if perspective == 'teachers' else 'class_ids', [])
        if not included:
            continue
        d = p['start'][0]; first = int(p['start'][1:])
        teacher_names = '/'.join(names['teachers'][x]['name'] for x in s['teacher_ids'])
        subject_name = names['subjects'][s['subject_id']]['name']
        label = f'{subject_name}\n{teacher_names}'
        if p.get('room_id'):
            label += '\n'+names['rooms'][p['room_id']]['name']
        for n in range(first, first+s['duration']):
            cells[(d,n)] = label
    return cells


def export_xlsx(project, placements):
    wb = Workbook(); summary = wb.active; summary.title='배정목록'
    summary.append(['회차','과목','교사','학급','학생그룹','요일·교시','교실'])
    sess = {x['session_id']: x for x in expand_sessions(project)}
    for p in placements:
        s = sess[p['session_id']]
        values=[p['session_id'],s['subject_id'],', '.join(s.get('teacher_ids',[])),', '.join(s.get('class_ids',[])),
                ', '.join(s.get('student_ids',[])),p['start'],p.get('room_id','')]
        summary.append([excel_safe(value) for value in values])
    for category, prefix in [('classes','학급'),('teachers','교사'),('rooms','교실'),('students','학생')]:
        for item in project.get(category, []):
            if category == 'students' and not any(item['id'] in x.get('student_ids',[]) for x in project['lessons']):
                continue
            ws = wb.create_sheet((prefix+'_'+item['id'])[:31]); ws.append([prefix, item.get('name', item['id'])])
            ws.append(['교시']+DAYS)
            cells = timetable_rows(project, placements, category, item['id'])
            for period in range(1, max(project['periods_by_day'].values())+1):
                ws.append([period]+[excel_safe(cells.get((d,period),'')) for d in DAYS])
            ws.column_dimensions['A'].width = 8
            for column in 'BCDEF': ws.column_dimensions[column].width = 22
            for row in ws.iter_rows(min_row=3):
                for cell in row: cell.alignment=Alignment(wrap_text=True, vertical='center')
            for row in range(3, ws.max_row+1): ws.row_dimensions[row].height=40
            for cell in ws[2]:
                cell.font=Font(bold=True, color='FFFFFF'); cell.fill=PatternFill('solid', fgColor='176B79')
    output = BytesIO(); wb.save(output); return output.getvalue()
