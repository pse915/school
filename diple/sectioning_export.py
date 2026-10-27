"""Administrator-only spreadsheet reporting of Phase 5B allocation results."""
from io import BytesIO
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def plan_to_excel(snapshot, plan):
    from .sectioning import validate_plan,SectioningError
    problems=validate_plan(snapshot,plan)
    if plan.get('resources_confirmed'):
        from .resources import resource_errors
        problems+=resource_errors(plan)
    if problems:raise SectioningError('검증 실패한 계획은 엑셀로 내보낼 수 없습니다.')
    wb=Workbook();meta=wb.active;meta.title='PlanVersion'
    for key in ('plan_id','version','source_snapshot_id','source_snapshot_version','school_id','academic_year','semester','campaign_id','created_at','solver','status'):
        meta.append([key,plan.get(key,'')])
    for warning in plan.get('warnings',[]):meta.append(['warning',warning])
    actual=wb.create_sheet('ActualTimetable')
    actual.append(['section_id','teacher_id','room_id','weekday_periods','resources_confirmed'])
    for sec in plan['sections']:
        actual.append([sec['section_id'],sec.get('teacher_id',''),sec.get('room_id',''),','.join(sec.get('slots',[])),bool(plan.get('resources_confirmed'))])
    timetable=wb.create_sheet('StudentTimetable')
    timetable.append(['student_id','course_name','weekday_period','teacher_id','room_id','resources_confirmed'])
    secmap={sec['section_id']:sec for sec in plan['sections']}
    cname={o['offering_id']:o['course_name'] for o in snapshot['offerings']}
    for assignment in plan['assignments']:
        sec=secmap[assignment['section_id']]
        for slot in sec.get('slots',[]):
            timetable.append([assignment['student_id'],cname[assignment['offering_id']],slot,sec.get('teacher_id',''),sec.get('room_id',''),bool(plan.get('resources_confirmed'))])
    groups=wb.create_sheet('OperatingBlocks')
    groups.append(['group_id','group_name','units_per_lane','lane_count','solver'])
    for g in plan['groups']:
        groups.append([g['group_id'],g['name'],g['lane_units'],g['lane_count'],g['solver']])
    items=wb.create_sheet('Sections')
    items.append(['section_id','offering_id','course_id','course_name','group_id','lane_1based','offset_1based','units','enrolled','capacity'])
    count=Counter(a['section_id'] for a in plan['assignments'])
    names={o['offering_id']:o['course_name'] for o in snapshot['offerings']}
    for s in plan['sections']:
        items.append([s['section_id'],s['offering_id'],s['course_id'],names[s['offering_id']],s['group_id'],
                      s['lane']+1,s['offset']+1,s['units'],count[s['section_id']],s['capacity']])
    roster=wb.create_sheet('StudentAssignment')
    roster.append(['snapshot_id','student_id','grade','class_id','number','offering_id','course_id','course_name','section_id','group_id','lane_1based','offset_1based','units'])
    students={s['student_id']:s for s in snapshot['students']};secs={s['section_id']:s for s in plan['sections']}
    for a in plan['assignments']:
        s=secs[a['section_id']];u=students[a['student_id']]
        roster.append([snapshot['snapshot_id'],a['student_id'],u['grade'],u['class_id'],u['number'],
                       a['offering_id'],s['course_id'],names[a['offering_id']],a['section_id'],s['group_id'],
                       s['lane']+1,s['offset']+1,s['units']])
    warning=wb.create_sheet('ImportValidation')
    warning.append(['attribute','instruction'])
    warning.append(['source_snapshot_id','반드시 원본 5A 확정 스냅샷 ID와 일치시켜야 합니다.'])
    warning.append(['student_id','안정적인 내부 학생 ID입니다. 학번으로 매칭하지 마세요.'])
    warning.append(['section_id','운영그룹·교과·구획·정원과 함께 검증해야 합니다.'])
    warning.append(['note','이 문서는 분반 검토 보고서입니다. 외부 소프트웨어 전용 호환 양식으로 인증되지 않았습니다.'])
    for sheet in (groups,items,roster,warning):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:
            c.font=Font(color='FFFFFF',bold=True)
            c.fill=PatternFill('solid',fgColor='245A75')
            c.alignment=Alignment(vertical='center',wrap_text=True)
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(46,max(14,max(len(str(c.value or '')) for c in list(col)[:55])+2))
    out=BytesIO();wb.save(out);return out.getvalue()
