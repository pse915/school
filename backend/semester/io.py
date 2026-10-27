from io import BytesIO
import csv,json,html
from copy import deepcopy
import pandas as pd
from openpyxl import Workbook,load_workbook
from diple.model import expand_sessions
from .model import SemesterError,validate

FIELDS={
 'teachers':['id','name','subject','max_weekly','blocked','working_weekdays','available_slots','starts_at','ends_at'],
 'classes':['id','name','grade','size','blocked'],
 'rooms':['id','name','capacity','type','blocked'],
 'subjects':['id','name','credits'],
 'students':['id','name','class_id'],
 'curriculum':['id','subject_id','teacher_ids','class_ids','student_ids','room_ids','weekly_periods','split','fixed_slots','forbidden_slots','simultaneous_group','parts'],
}
LISTS={'blocked','working_weekdays','available_slots','teacher_ids','class_ids','student_ids','room_ids','split','fixed_slots','forbidden_slots','parts'}
NUMBERS={'max_weekly','grade','size','capacity','credits','weekly_periods'}

def rows_for_editor(project,category):
    return [{key:json.dumps(row[key],ensure_ascii=False) if key in LISTS and key in row else row.get(key,'') for key in FIELDS[category]} for row in project[category]]

def parse_rows(rows,category):
    result=[];seen=set()
    for i,row in enumerate(rows,1):
        if all(value is None or str(value).strip() in ('','nan') for value in row.values()):continue
        item={}
        try:
            for key in FIELDS[category]:
                value=row.get(key)
                if value is None or str(value).strip() in ('','nan'):continue
                if key in LISTS:
                    value=value if isinstance(value,list) else json.loads(str(value)) if str(value).strip().startswith('[') else str(value).replace('+',',').split(',')
                    if key=='split':value=[int(x) for x in value]
                    elif key!='parts':value=[str(x).strip() for x in value if str(x).strip()]
                elif key in NUMBERS:
                    if float(value)!=int(float(value)):raise ValueError('정수가 필요합니다.')
                    value=int(float(value))
                else:value=str(value).strip()
                item[key]=value
            ident=item.get('id','')
            if not ident or ident in seen:raise ValueError('비어 있거나 중복된 ID')
            seen.add(ident)
            if category=='curriculum':
                for key in ('class_ids','student_ids','room_ids','forbidden_slots','fixed_slots'):item.setdefault(key,[])
            if category in ('teachers','classes','rooms'):item.setdefault('blocked',[])
        except (ValueError,TypeError) as exc:raise SemesterError(f'{category} {i}행: {exc}') from exc
        result.append(item)
    return result

def template():
    wb=Workbook();wb.remove(wb.active)
    for category,fields in FIELDS.items():
        ws=wb.create_sheet(category);ws.append(fields);ws.freeze_panes='A2'
    meta=wb.create_sheet('사용방법');meta.append(['목록은 JSON 배열 또는 쉼표 구분. split은 2+1도 가능. 학점과 주당 시수는 별도.'])
    meta.append(['교사 분담: 같은 과목의 별도 배당행 또는 parts=[{"duration":2,"teacher_ids":["T1"]},{"duration":1,"teacher_ids":["T2"]}]. 공동수업: teacher_ids에 두 교사.'])
    out=BytesIO();wb.save(out);return out.getvalue()

def read_upload(upload,sheet):
    upload.seek(0)
    if upload.name.lower().endswith('.csv'):return pd.read_csv(upload,dtype=object).fillna('')
    return pd.read_excel(upload,sheet_name=sheet,dtype=object).fillna('')

def timetable_rows(project,result):
    sessions={s['session_id']:s for s in expand_sessions(project)};teachers={t['id']:t['name'] for t in project['teachers']};subjects={s['id']:s['name'] for s in project['subjects']}
    return [{'회차ID':p['session_id'],'배당ID':sessions[p['session_id']]['source_request_id'],'요일':p['start'][0],'시작교시':int(p['start'][1:]),'연속시수':sessions[p['session_id']]['duration'],
             '과목':subjects[sessions[p['session_id']]['subject_id']],'교사ID':','.join(sessions[p['session_id']]['teacher_ids']),
             '교사':','.join(teachers[t] for t in sessions[p['session_id']]['teacher_ids']),'학급':','.join(sessions[p['session_id']]['class_ids']),
             '학생ID':','.join(sessions[p['session_id']]['student_ids']),'교실':p.get('room_id','')}
            for p in sorted(result['placements'],key=lambda x:(x['start'][0],int(x['start'][1:]),x['session_id']))]

def export_excel(project,result,metadata=None):
    errors=validate(project,result['placements'])
    if errors:raise SemesterError('검증 오류가 있는 시간표는 확정 출력할 수 없습니다.')
    rows=timetable_rows(project,result);wb=Workbook();meta=wb.active;meta.title='버전정보'
    for k,v in {**project['school'],**(metadata or {}),'status':result['status']}.items():meta.append([k,str(v)])
    for title,key in [('전체',''),('학급별','학급'),('교사별','교사ID'),('교실별','교실'),('이동반별','배당ID')]:
        ws=wb.create_sheet(title);ws.append(list(rows[0]) if rows else ['자료 없음'])
        for row in sorted(rows,key=lambda r:str(r.get(key,''))):ws.append(list(row.values()))
    student=wb.create_sheet('학생별');student.append(['학생ID','소속학급','과목','요일','시작교시','연속시수','교사','교실'])
    sessions={s['session_id']:s for s in expand_sessions(project)}
    for s in project['students']:
        for row in rows:
            lesson=sessions[row['회차ID']]
            if s['id'] in lesson['student_ids'] or s['class_id'] in lesson['class_ids']:student.append([s['id'],s['class_id'],row['과목'],row['요일'],row['시작교시'],row['연속시수'],row['교사'],row['교실']])
    summary=wb.create_sheet('교사배당시수');summary.append(['교사ID','교사명','주당시수'])
    for teacher in project['teachers']:summary.append([teacher['id'],teacher['name'],sum(s['duration'] for s in sessions.values() if teacher['id'] in s['teacher_ids'])])
    report=wb.create_sheet('검증');report.append(['전체 검증 오류',len(errors)])
    for k,v in result.get('metrics',{}).items():report.append([k,v])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
    out=BytesIO();wb.save(out);return out.getvalue()

def export_html(project,result):
    errors=validate(project,result['placements'])
    if errors:raise SemesterError('검증 오류가 있는 시간표는 인쇄할 수 없습니다.')
    rows=timetable_rows(project,result);keys=['요일','시작교시','연속시수','과목','교사','학급','교실']
    text='<!doctype html><html lang="ko"><meta charset="utf-8"><title>학기 시간표</title><style>body{font-family:sans-serif}table{border-collapse:collapse;width:100%}td,th{border:1px solid #999;padding:6px}@media print{button{display:none}}</style><h1>'+html.escape(project['term']['name'])+'</h1><button onclick="print()">인쇄 / PDF 저장</button><table><thead><tr>'+''.join('<th>'+k+'</th>' for k in keys)+'</tr></thead><tbody>'
    text+=''.join('<tr>'+''.join('<td>'+html.escape(str(row[k]))+'</td>' for k in keys)+'</tr>' for row in rows)+'</tbody></table></html>'
    return text.encode('utf-8')
